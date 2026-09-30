"""Verify Foundry manuscript, native PPTX objects, notes, and round-trip provenance."""
import hashlib
import base64
import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

from lxml import etree, html
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def compact(text):
    return re.sub(r"\s+", "", text)


def digest(file):
    return hashlib.sha256(file.read_bytes()).hexdigest()


def main():
    catalog = json.loads((ROOT / "catalog.json").read_text())
    source_map = json.loads((ROOT / "shared/sources.json").read_text())["sources"]
    markdown = (ROOT / "shared/foundry.md").read_text()
    sections = markdown.split("\n## ")[1:]
    specs = [json.loads(re.search(r"```json\n([\s\S]*?)\n```", s)[1]) for s in sections]
    assert len(specs) == 15
    layouts = [spec["kind"] for spec in specs]
    results = []
    for entry in catalog["samples"]:
        assert entry["slides"] == 15 and entry["layouts"] == layouts, entry["skill"]
        coverage = entry["coverage"]
        assert digest(ROOT.parent / coverage["template"]) == coverage["template_sha256"]
        assert [s["layout"] for s in coverage["slides"]] == layouts
        if entry["output"].endswith(".html"):
            document = html.fromstring((ROOT / entry["output"]).read_bytes(), parser=html.HTMLParser(huge_tree=True))
            assert len(document.xpath('//*[@class="chapter"] | //article[contains(concat(" ",@class," ")," chapter ")]')) == 15
            for spec in specs:
                assert compact(spec["title"]) in compact(document.text_content()), (entry["skill"], spec["kind"])
            assert document.xpath('//script[@id="standalone-provenance"]')
            assert document.xpath('//script[@id="bundled-font-licenses"]')
        if not entry["output"].endswith(".pptx"):
            continue
        file = ROOT / entry["output"]
        deck = Presentation(file)
        assert len(deck.slides) == len(specs), file
        timing_count = 0
        manifest_file = file.parent / "anim-manifest.json"
        manifest = json.loads(manifest_file.read_text()) if manifest_file.exists() else None
        if manifest:
            for effect in (effect for row in manifest for effect in row if effect):
                assert abs(effect["off"] - 30 / 1080) < 0.00001, (file, "entrance distance")
                duration = 620 if entry["skill"].startswith("deep-navy") else 700
                assert effect["dur"] == duration, (file, "entrance duration")
                assert effect["d"] > 0, (file, "lost stagger delay")
        with ZipFile(file) as archive:
            for index, (slide, spec, section) in enumerate(zip(deck.slides, specs, sections), 1):
                text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                text += "\n".join(cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
                notes = slide.notes_slide.notes_text_frame.text.replace("\v", "\n")
                all_text = compact(text + "\n" + notes)
                assert compact(spec["title"]) in compact(text), (file, index, "headline")
                for field in ("lead", "context", "insight", "caution"):
                    assert compact(spec[field]) in compact(text), (file, index, field)
                for item in spec.get("items", []):
                    for value in item.values():
                        assert compact(value) in all_text, (file, index, value)
                for source in spec["sources"]:
                    assert source_map[source]["url"] in notes, (file, index, "source URL")
                narrative = section.split("```\n", 2)[-1].strip().replace("### ", "")
                assert compact(narrative) in compact(notes), (file, index, "narrative/notes")
                assert "\u2424" not in notes, (file, index, "unrestored notes")
                if "rows" in spec:
                    tables = [shape.table for shape in slide.shapes if shape.has_table]
                    assert len(tables) == 1, (file, index, "native table")
                    actual = [[cell.text for cell in row.cells] for row in tables[0].rows]
                    assert actual == [spec["head"], *spec["rows"]], (file, index, "table values")
                for shape in slide.shapes:
                    assert shape.left >= 0 and shape.top >= 0, (file, index, "negative position")
                    assert shape.left + shape.width <= deck.slide_width + 100, (file, index, "width overflow")
                    assert shape.top + shape.height <= deck.slide_height + 100, (file, index, "height overflow")
                xml = etree.fromstring(archive.read(f"ppt/slides/slide{index}.xml"))
                effects = xml.xpath("//p:animEffect", namespaces=NS)
                timing_count += len(effects)
                ids = [node.get("id") for node in xml.xpath("//p:cNvPr", namespaces=NS)]
                assert len(ids) == len(set(ids)), (file, index, "duplicate ID")
                if manifest:
                    expected = sum(bool(effect) for effect in manifest[index - 1])
                    assert len(effects) == expected, (file, index, "animation count")
                else:
                    assert not effects, (file, index, "unexpected animation")
        pdf_text = subprocess.check_output(["pdftotext", "-layout", str(file.with_suffix(".pdf")), "-"], text=True)
        pdf_pages = pdf_text.rstrip("\f\n").split("\f")
        assert len(pdf_pages) == len(specs), (file, "PDF page count")
        for spec, page in zip(specs, pdf_pages):
            # Layout extraction interleaves text from parallel columns.
            for line in spec["title"].splitlines():
                assert compact(line) in compact(page), (file, "PDF headline line missing", line)
        results.append(dict(file=entry["output"], slides=len(deck.slides), animations=timing_count,
                            native_tables=sum(shape.has_table for slide in deck.slides for shape in slide.shapes),
                            notes_and_sources="preserved", pdf_headlines="present"))
    for theme in ("deep-navy", "white-cobalt"):
        folder = ROOT / ".cache/work" / f"{theme}-pptx-to-html"
        manifest = json.loads((folder / "manifest.json").read_text())
        document = html.fromstring((ROOT / f"{theme}-pptx-to-html/index.html").read_bytes(), parser=html.HTMLParser(huge_tree=True))
        embedded = json.loads(document.xpath('//script[@id="source-preservation"]')[0].text)
        assert embedded == manifest
        downloads = {a.get("download"): base64.b64decode(a.get("href").split(",", 1)[1])
                     for a in document.xpath('//a[@download]')}
        assert hashlib.sha256(downloads["source.pptx"]).hexdigest() == manifest["source_sha256"]
        assert hashlib.sha256(downloads["source.pdf"]).hexdigest() == manifest["pdf_sha256"]
        assert digest(folder / "source.pptx") == manifest["source_sha256"]
        assert digest(folder / "source.pdf") == manifest["pdf_sha256"]
        assert digest(folder / "source.pptx") == digest(ROOT / f"{theme}-html-to-pptx/foundry.pptx")
        assert manifest["slides"] == len(specs)
        for page, spec in zip(manifest["pages"], specs):
            assert compact(page["title"]) == compact(spec["title"])
            assert digest(folder / page["png"]) == page["png_sha256"]
            assert digest(folder / page["svg"]) == page["sha256"]
            for source in spec["sources"]:
                assert source_map[source]["url"] in page["notes"]
        for asset in manifest["assets"]:
            assert digest(folder / asset["file"]) == asset["sha256"]
            assert hashlib.sha256(downloads[Path(asset["file"]).name]).hexdigest() == asset["sha256"]
    for font in json.loads((ROOT / "assets/fonts/provenance.json").read_text()):
        assert digest(ROOT / "assets/fonts" / font["file"]) == font["sha256"]
    (ROOT / "content-validation.json").write_text(json.dumps({
        "template_coverage": {"samples": len(catalog["samples"]), "templates_per_sample": 15,
                              "source_template_hashes": "match"},
        "pptx": results, "source_preservation": "hashes/pages/notes/titles match",
        "font_provenance": "all hashes match", "powerpoint_playback": "not verified",
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("PPTX content, tables, notes, source URLs, bounds, timing, PDF text and round-trip provenance: OK")


if __name__ == "__main__":
    main()
