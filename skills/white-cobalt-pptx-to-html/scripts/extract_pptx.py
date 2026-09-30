#!/usr/bin/env python3
"""Extract editable PPTX content; optionally scaffold either reading theme."""
import argparse
import html
import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build_html import THEMES, assemble, template_path


def extract(source, image_dir):
    deck = Presentation(source)
    slides = []
    for index, slide in enumerate(deck.slides, 1):
        texts, tables, images, warnings = [], [], [], []

        def visit(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    embeds = shape._element.xpath(".//a:blip/@r:embed")
                    svg_embeds = shape._element.xpath('.//*[local-name()="svgBlip"]/@r:embed')
                    if not embeds and not svg_embeds:
                        warnings.append(f"Picture '{shape.name}' has no embedded image; inspect the source for linked media.")
                        continue
                    image_dir.mkdir(parents=True, exist_ok=True)
                    # An SVG-only picture is still an embedded original, not missing media.
                    part = slide.part.related_part((svg_embeds or embeds)[0])
                    suffix = Path(str(part.partname)).suffix
                    out = image_dir / f"s{index:02d}-{len(images) + 1}{suffix}"
                    out.write_bytes(part.blob)
                    images.append(str(out.resolve()))
                elif shape.has_table:
                    tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
                elif shape.has_chart:
                    warnings.append(f"Chart '{shape.name}' requires original-slide visual comparison.")
                elif shape.has_text_frame and shape.text.strip():
                    texts.append(shape.text.strip().replace("\v", "\n"))
                else:
                    warnings.append(f"Visual shape '{shape.name}' needs manual comparison with the original.")
        visit(slide.shapes)
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else (texts[0] if texts else f"Slide {index}")
        if title in texts:
            texts.remove(title)
        notes = slide.notes_slide.notes_text_frame.text.replace("\v", "\n") if slide.has_notes_slide else ""
        slides.append(dict(n=index, title=title, texts=texts, tables=tables, images=images, notes=notes, warnings=warnings))
    return slides


def fragments(slides, theme, output_dir):
    """Scaffold only: full extracted content stays in the reader explanation."""
    chapters = []
    esc = html.escape
    for item in slides:
        title = esc(item["title"])
        snippets = [text.splitlines()[0] for text in item["texts"] if len(text.splitlines()[0]) <= 64][:4]
        if theme == "deep-navy":
            body = '<div class="frame"><header class="s-head"><h2 class="s-title">' + title + "</h2></header>"
            body += '<div class="s-body"><div class="list tight">'
            for number, text in enumerate(snippets, 1):
                body += f'<div class="li"><div class="b">{number:02d}</div><div class="tx"><div class="t">{esc(text)}</div></div></div>'
            body += '</div></div><div class="s-foot"><span>원본 요약</span><span class="pg"></span></div></div>'
        else:
            body = '<header class="topbar"><div class="sec">원본 요약</div><div class="pg"></div></header><div class="body">'
            body += f'<h2>{title}</h2><ul class="after" style="margin-top:40px">'
            body += "".join(f"<li>{esc(text)}</li>" for text in snippets) + "</ul></div>"
        explanation = "".join("<p>" + esc(text).replace("\n", "<br>") + "</p>" for text in item["texts"])
        for table in item["tables"]:
            explanation += '<table><tbody>' + "".join("<tr>" + "".join(f"<td>{esc(cell)}</td>" for cell in row) + "</tr>" for row in table) + "</tbody></table>"
        for image in item["images"]:
            relative = Path(os.path.relpath(image, output_dir)).as_posix()
            explanation += f'<p><img src="{esc(relative, quote=True)}" alt="슬라이드 {item["n"]} 원본 이미지"></p>'
        explanation += "<p>" + esc(item["notes"]).replace("\n", "<br>") + "</p>"
        explanation += "".join(f'<p><b>변환 주의:</b> {esc(w)}</p>' for w in item["warnings"])
        chapters.append(
            f'<article class="chapter" data-part="원본 자료"><div class="ch-head"><span class="ch-n"></span>'
            f'<h2 class="ch-t">{title}</h2><span class="ch-part">원본 자료</span><a class="ch-link" href="#">#</a></div>'
            f'<div class="sheet"><div class="slide">{body}</div></div>'
            f'<aside class="note"><div class="note-head">슬라이드 설명</div><div class="note-body">{explanation}</div></aside></article>')
    return "\n".join(chapters)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("--img-dir", type=Path, default=Path("images"))
    parser.add_argument("--json", type=Path, required=True)
    parser.add_argument("--fragments", type=Path)
    parser.add_argument("--html", type=Path, help="Optional reflow scaffold, not an original-slide preservation")
    parser.add_argument("--theme", choices=THEMES, default="white-cobalt")
    args = parser.parse_args()
    slides = extract(args.pptx, args.img_dir)
    args.json.parent.mkdir(parents=True, exist_ok=True)
    args.json.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.fragments:
        args.fragments.parent.mkdir(parents=True, exist_ok=True)
        args.fragments.write_text(fragments(slides, args.theme, args.fragments.parent.resolve()), encoding="utf-8")
        print("Scaffold only: review the original slide layout and rewrite coaching notes before delivery.")
    if args.html:
        args.html.parent.mkdir(parents=True, exist_ok=True)
        args.html.write_text(assemble(
            fragments(slides, args.theme, args.html.parent.resolve()),
            template_path(args.theme).read_text(encoding="utf-8"), Path(args.pptx).stem,
            subtitle="Reflow scaffold — inspect against the original presentation"),
            encoding="utf-8")
        print("HTML scaffold only: source rendering, layout and animations are not preserved.")
    for slide in slides:
        for warning in slide["warnings"]:
            print(f'Slide {slide["n"]}: {warning}')
    print(f"{len(slides)} slides extracted to {args.json}")


if __name__ == "__main__":
    main()
