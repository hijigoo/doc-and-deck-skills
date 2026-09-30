#!/usr/bin/env python3
"""Wrap original slide renders in a reading UI; never summarize source content."""
import argparse
import base64
import hashlib
import html
import io
import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from build_html import THEMES, assemble, template_path

XLINK = "{http://www.w3.org/1999/xlink}href"


def digest(data):
    return hashlib.sha256(data).hexdigest()


def pixel_key(data):
    with Image.open(io.BytesIO(data)) as image:
        # Avoid changing ICC/gamma interpretation or making a static GIF animate.
        if image.info.get("icc_profile") or getattr(image, "n_frames", 1) != 1:
            return None
        rgba = image.convert("RGBA")
        return (rgba.size, digest(rgba.tobytes()))


def raw_text(shapes):
    parts = []
    for shape in shapes:
        if shape.shape_type == 6:
            parts.extend(raw_text(shape.shapes))
        elif shape.has_text_frame and shape.text:
            parts.append(shape.text)
        elif shape.has_table:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return parts


def normalize_rendered_pages(pages, count):
    """Poppler pads to the PDF page count, not consistently to two digits."""
    rendered = {}
    for path in pages.glob("rendered-*.png"):
        match = re.fullmatch(r"rendered-(\d+)\.png", path.name)
        if match:
            index = int(match[1])
            if index in rendered:
                raise ValueError(f"Duplicate raster render for slide {index}.")
            rendered[index] = path
    if set(rendered) != set(range(1, count + 1)):
        raise ValueError("Rasterized PDF pages do not match the source slide count.")
    for index, path in rendered.items():
        path.replace(pages / f"slide-{index:02d}.png")


def preserve(source, pdf, output, renderer, theme="deep-navy", template=None):
    source, pdf, output = Path(source).resolve(), Path(pdf).resolve(), Path(output).resolve()
    template = Path(template) if template is not None else template_path(theme)
    template_text = template.read_text(encoding="utf-8")
    presentation = Presentation(str(source))
    count = len(presentation.slides)
    if not count:
        raise ValueError("The source presentation contains no slides.")
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    if int(re.search(r"^Pages:\s+(\d+)", info, re.M)[1]) != count:
        raise ValueError("PDF/PPTX page count mismatch. Include hidden slides in a working render copy.")
    assets, pages = output / "assets", output / "pages"
    assets.mkdir(parents=True, exist_ok=True)
    pages.mkdir(exist_ok=True)
    # Preserve the native PDF rasterization as the primary view. SVG remains
    # available for zoom, but browser vector anti-aliasing can differ visibly.
    subprocess.run(["pdftoppm", "-png", "-r", "144", str(pdf), str(pages / "rendered")], check=True)
    normalize_rendered_pages(pages, count)
    manifest, pixels = [], {}
    with ZipFile(source) as archive:
        for entry in archive.infolist():
            if not entry.filename.startswith("ppt/media/") or entry.is_dir():
                continue
            data = archive.read(entry)
            target = assets / Path(entry.filename).name
            target.write_bytes(data)
            manifest.append(dict(part=entry.filename, file=str(target.relative_to(output)),
                                 sha256=digest(data), bytes=len(data)))
            if target.suffix.lower() in (".png", ".gif", ".jpg", ".jpeg"):
                key = pixel_key(data)
                if key:
                    pixels[key] = target
    chapters, slides, direct_references = [], [], 0
    ratio = presentation.slide_width / presentation.slide_height
    for n, slide in enumerate(presentation.slides, 1):
        target = pages / f"slide-{n:02d}.svg"
        subprocess.run(["pdftocairo", "-svg", "-f", str(n), "-l", str(n), str(pdf), str(target)], check=True)
        xml = etree.parse(str(target))
        used = []
        for image in xml.xpath('//*[local-name()="image"]'):
            uri = image.get(XLINK, image.get("href", ""))
            if uri.startswith("data:image/png;base64,"):
                key = pixel_key(base64.b64decode(uri.split(",", 1)[1]))
                if key in pixels:
                    original = pixels[key]
                    mime = "image/jpeg" if original.suffix.lower() in (".jpg", ".jpeg") else "image/" + original.suffix[1:].lower()
                    # SVGs used through <img> cannot load external image URLs.
                    # Embed the identical original bytes, not a newly encoded bitmap.
                    image.set(XLINK, f"data:{mime};base64," + base64.b64encode(original.read_bytes()).decode())
                    used.append(str(original.relative_to(output)))
        direct_references += len(used)
        xml.write(str(target), encoding="UTF-8", xml_declaration=True)
        texts = raw_text(slide.shapes)
        heading = slide.shapes.title.text if slide.shapes.title is not None else ""
        if not heading.strip():
            heading = next((text for text in texts if text.strip() and not text.strip().isdigit()), f"원본 슬라이드 {n}")
        heading = re.sub(r"\s+", " ", heading).strip()
        hidden = slide._element.get("show") == "0"
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        shown_title = html.escape(heading)
        state = " · 원본 숨김 장표" if hidden else ""
        transcript = "".join("<p>" + html.escape(text) + "</p>" for text in texts)
        note_html = html.escape(notes).replace("\v", "\n") if notes else "원본 발표자 노트가 없습니다."
        links = [str(rel.target_ref) for rel in slide.part.rels.values() if rel.reltype == RT.HYPERLINK]
        link_html = ""
        if links:
            link_html = "<h3>원본 하이퍼링크</h3><ul>" + "".join(
                '<li><a href="' + html.escape(link, quote=True) + '">' + html.escape(link) + "</a></li>"
                if link.startswith(("https://", "http://", "mailto:")) else "<li>" + html.escape(link) + "</li>"
                for link in links) + "</ul>"
        chapters.append(
            f'<article class="chapter" data-part="원본 슬라이드" data-source-slide="{n}" data-source-hidden="{str(hidden).lower()}">'
            f'<div class="ch-head"><span class="ch-n"></span><h2 class="ch-t">{shown_title}</h2><span class="ch-part">원본 {n}{state}</span><a class="ch-link" href="#">#</a></div>'
            f'<div class="sheet"><div class="slide preserved-slide" data-source-ratio="{ratio}"><a href="pages/slide-{n:02d}.svg" target="_blank" rel="noopener" aria-label="원본 {n}장 크게 보기">'
            f'<img src="pages/slide-{n:02d}.png" alt="원본 {n}장: {html.escape(heading,quote=True)}"></a></div></div>'
            '<aside class="note"><div class="note-head">원본 노트와 텍스트</div><div class="note-body">'
            f'<pre class="source-notes">{note_html}</pre>{link_html}'
            f'<details class="source-transcript"><summary>원문 텍스트 복사</summary>{transcript}</details>'
            '</div></aside></article>')
        slides.append(dict(n=n, title=heading, hidden=hidden, text=texts, notes=notes, links=links,
                           svg=str(target.relative_to(output)), sha256=digest(target.read_bytes()),
                           png=f"pages/slide-{n:02d}.png",
                           png_sha256=digest((pages / f"slide-{n:02d}.png").read_bytes()),
                           copied_asset_references=used))
    hidden = [s["n"] for s in slides if s["hidden"]]
    title = source.stem
    doc = assemble("\n".join(chapters), template_text, title, subtitle=f"원본 {count}장 · 내용과 배치 보존",
                   kicker="PPTX → WEB",
                   intro=f"원본 슬라이드를 요약·재작성하지 않고 본문에 그대로 표시합니다. 원본 자산은 별도 파일로 복사했습니다. 시트는 {renderer} PDF에서 만든 고해상도 이미지이며 클릭하면 확대용 SVG가 열립니다. 글자 선택은 아래 원문 텍스트를 이용합니다. 숨김 장표도 원래 순서에 포함하고 표시했습니다: {hidden or '없음'}.",
                   meta=f"{count} Original slides|No summarization|Native render")
    extra = f"""
    .sheet{{aspect-ratio:{ratio}}}
    .sheet .preserved-slide{{padding:0!important;width:1920px;height:{1920/ratio}px;display:block}}
    .preserved-slide a,.preserved-slide img{{display:block;width:100%;height:100%;margin:0;max-width:none;max-height:none}}
    .preserved-slide img{{object-fit:contain}}
    .source-notes,.source-transcript p{{white-space:pre-wrap;overflow-wrap:anywhere;font-family:var(--f-body);font-size:16px;line-height:1.7}}
    .source-notes{{padding:8px 0}}.source-transcript{{margin-top:18px}}.source-transcript summary{{cursor:pointer;color:var(--cobalt);font-weight:bold}}
    .source-transcript p{{border-top:1px solid var(--line);padding-top:10px;margin-top:10px}}
    """
    doc = doc.replace("</style>", extra + "\n</style>")
    (output / "index.html").write_text(doc, encoding="utf-8")
    result = dict(source=source.name, source_sha256=digest(source.read_bytes()), pdf_sha256=digest(pdf.read_bytes()),
                  renderer=renderer, slides=count, hidden_slides=hidden, aspect_ratio=ratio,
                  assets=manifest, pages=slides, original_asset_references_in_svg=direct_references,
                  limitations=["Static source render; original animations/video and editable slide text are not reproduced.",
                               "Original link destinations are preserved in notes, not as exact on-slide click regions.",
                               "PDF rendering fidelity depends on the export application and installed fonts.",
                               "Original media files are copied byte-for-byte; some displayed graphics are source-derived PDF vectors."])
    (output / "manifest.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{count} original slides; {len(manifest)} original media files; {direct_references} direct asset references in SVG.")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("--pdf", required=True, help="PDF export of every source slide in its original order")
    parser.add_argument("--out", required=True)
    parser.add_argument("--renderer", required=True, help="Actual PDF export application; never claim a different renderer")
    parser.add_argument("--theme", choices=THEMES, default="deep-navy")
    parser.add_argument("--template", type=Path, help="Optional reading template override")
    args = parser.parse_args()
    preserve(args.pptx, args.pdf, args.out, args.renderer, args.theme, args.template)
